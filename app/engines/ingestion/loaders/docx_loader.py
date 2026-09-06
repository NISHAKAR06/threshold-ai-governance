"""
docx_loader.py — DOCX document loader using python-docx.
"""
from pathlib import Path
from typing import Union
from docx import Document
from app.engines.ingestion.loaders.base_loader import BaseLoader
from app.core.exceptions import LoaderError


class DOCXLoader(BaseLoader):
    """Extracts text content from Microsoft Word (.docx) documents, including paragraphs and tables."""

    def load(self, file_path: Union[str, Path]) -> str:
        p = Path(file_path)
        if not p.exists():
            raise LoaderError(f"DOCX file does not exist: {p}", str(p))

        try:
            doc = Document(str(p))
            parts = []

            # Extract tables first if any (e.g. document headers)
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    parts.append("\n".join(table_lines))

            # Extract paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)

            full_text = "\n\n".join(parts)
            if not full_text.strip():
                raise LoaderError(f"DOCX file contains no readable text: {p}", str(p))

            return full_text
        except Exception as e:
            if isinstance(e, LoaderError):
                raise e
            raise LoaderError(f"Failed to extract text from DOCX '{p}': {str(e)}", str(p))
